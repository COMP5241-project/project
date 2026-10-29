import ipaddress
import logging
import os
import re
import socket
import uuid
from collections.abc import Generator
from datetime import UTC, datetime
from html.parser import HTMLParser
from typing import Any, Literal
from urllib.parse import urljoin, urlsplit

from fastapi import APIRouter, Depends, Request, Response
from httpx import Client as HttpClient
from pydantic import BaseModel, ConfigDict, Field, HttpUrl
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from starlette.datastructures import UploadFile as StarletteUploadFile
from supabase import Client

from app.identity.auth import OptionalClientDependency
from app.identity.errors import ApiError
from app.identity.profiles import load_profile
from app.identity.sessions import StoreDependency, verified_session
from app.pptx_export import PPTX_MEDIA_TYPE, build_pptx
from app.routers.auth import require_trusted_origin

router = APIRouter(prefix="/api/materials")
logger = logging.getLogger(__name__)
MAX_SOURCE_BYTES = 5 * 1024 * 1024
MAX_SOURCE_TEXT = 100_000
BUCKET_NAME = "instructor-materials"
# Descriptive, crawler-convention User-Agent. Some sites (e.g. Wikipedia)
# reject requests from server IPs unless the UA follows the
# ``Name/Version (contact-url; contact-email)`` convention, so override
# SOURCE_FETCH_USER_AGENT with the deployment's real contact details.
SOURCE_FETCH_USER_AGENT = os.getenv(
    "SOURCE_FETCH_USER_AGENT",
    "EduGenAI-SourceFetcher/1.0 (https://example.org; sources@example.org) python-httpx",
)
SourceType = Literal["file", "url", "text"]
MaterialType = Literal["slides", "quiz", "flashcards", "summary"]


class SourceItem(BaseModel):
    id: str
    material_id: str = Field(alias="materialId")
    type: SourceType
    name: str
    raw_content: str | None = Field(default=None, alias="rawContent")
    file_url: str | None = Field(default=None, alias="fileUrl")
    source_url: str | None = Field(default=None, alias="sourceUrl")
    created_at: str = Field(alias="createdAt")

    model_config = ConfigDict(populate_by_name=True)


class MaterialRecord(BaseModel):
    id: str
    user_id: str = Field(alias="userId")
    title: str
    type: MaterialType
    sources: list[SourceItem]
    data: dict[str, Any]
    chat_history: list[dict[str, Any]] = Field(alias="chatHistory")
    published: bool = False
    published_at: str | None = Field(default=None, alias="publishedAt")
    updated_at: str = Field(alias="updatedAt")
    created_at: str = Field(alias="createdAt")

    model_config = ConfigDict(populate_by_name=True)


class MaterialCreate(BaseModel):
    id: uuid.UUID | None = None
    title: str = Field(min_length=1, max_length=200)
    type: MaterialType
    data: dict[str, Any] = Field(default_factory=dict)
    chat_history: list[dict[str, Any]] = Field(default_factory=list, alias="chatHistory")

    model_config = ConfigDict(populate_by_name=True)


class MaterialPatch(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=200)
    data: dict[str, Any] | None = None
    chat_history: list[dict[str, Any]] | None = Field(default=None, alias="chatHistory")

    model_config = ConfigDict(populate_by_name=True)


class SourceCreate(BaseModel):
    type: Literal["text", "url"]
    name: str = Field(min_length=1, max_length=255)
    raw_content: str | None = Field(default=None, alias="rawContent", max_length=MAX_SOURCE_TEXT)
    source_url: HttpUrl | None = Field(default=None, alias="sourceUrl")

    model_config = ConfigDict(populate_by_name=True)


class TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "svg"}:
            self.ignored_depth += 1
        elif not self.ignored_depth and tag in {"p", "br", "div", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self.ignored_depth:
            self.ignored_depth -= 1
        elif not self.ignored_depth and tag in {"p", "div", "li", "h1", "h2", "h3", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            self.parts.append(data)


def _authorize_instructor(
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
) -> Generator[tuple[str, Client], None, None]:
    with verified_session(request, client, store) as (session, user):
        assert client is not None
        profile = load_profile(client, user.id, user.email, session.access_token).profile
        if profile.role != "instructor":
            raise ApiError(403, "instructor_required", "Only instructors can manage materials.")
        client.options.headers["Authorization"] = f"Bearer {session.access_token}"
        client.postgrest.auth(session.access_token)
        yield user.id, client


def _database_error(error: Exception) -> ApiError:
    logger.exception("Supabase material storage operation failed.")
    if getattr(error, "code", None) == "PGRST205":
        return ApiError(
            503,
            "materials_not_configured",
            "Instructor materials storage is not configured. Apply "
            "supabase/migrations/20261015000000_create_instructor_materials.sql "
            "to your Supabase project, then try again.",
        )
    if type(error).__name__ == "StorageApiError":
        status = str(getattr(error, "status", ""))
        code = str(getattr(error, "code", "")).lower()
        if status == "404" or code in {"bucketnotfound", "bucket not found"}:
            return ApiError(
                503,
                "source_bucket_missing",
                "The Supabase instructor-materials storage bucket is missing. Run "
                "supabase/migrations/20261015164500_configure_instructor_material_storage.sql "
                "in the Supabase SQL Editor to create the private bucket and its access policy.",
            )
        if status in {"400", "401", "403"}:
            return ApiError(
                503,
                "source_storage_forbidden",
                "Supabase denied the file upload. Confirm the signed-in user JWT is used for "
                "Storage requests and that the instructor-materials access policy is installed.",
            )
        return ApiError(
            503,
            "source_storage_not_configured",
            "Supabase file storage is not configured. Apply the complete instructor "
            "materials migration, including its private storage bucket and access policy.",
        )
    return ApiError(503, "materials_unavailable", "Unable to save or load instructor materials.")


def _material_row(
    row: dict[str, Any],
    sources: list[dict[str, Any]],
    client: Client,
) -> MaterialRecord:
    items: list[SourceItem] = []
    for source in sources:
        file_url = None
        storage_path = source.get("storage_path")
        if storage_path:
            try:
                result = client.storage.from_(BUCKET_NAME).create_signed_url(storage_path, 3600)
                file_url = result.get("signedURL") or result.get("signedUrl")
                if not file_url:
                    raise ValueError("Signed storage URL was not returned.")
            except Exception:
                logger.exception("Unable to create signed URL for an instructor source.")
                raise ApiError(
                    503, "source_storage_unavailable", "Unable to access a stored source file."
                ) from None
        items.append(SourceItem(
            id=str(source["id"]),
            materialId=str(source["material_id"]),
            type=source["type"],
            name=source["name"],
            rawContent=source.get("raw_content"),
            fileUrl=file_url,
            sourceUrl=source.get("source_url"),
            createdAt=source["created_at"],
        ))
    return MaterialRecord(
        id=str(row["id"]),
        userId=str(row["user_id"]),
        title=row["title"],
        type=row["type"],
        sources=items,
        data=row.get("data") or {},
        chatHistory=row.get("chat_history") or [],
        published=row.get("published", False),
        publishedAt=row.get("published_at"),
        updatedAt=row["updated_at"],
        createdAt=row["created_at"],
    )


def _owned_material(client: Client, material_id: str, user_id: str) -> dict[str, Any]:
    try:
        response = client.table("materials").select("*").eq("id", material_id).eq(
            "user_id", user_id
        ).maybe_single().execute()
    except Exception as error:
        raise _database_error(error) from None
    if not response or not response.data:
        raise ApiError(404, "material_not_found", "This material was not found.")
    return response.data


def _extract_file_text(filename: str, content_type: str, data: bytes) -> str:
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if not (
        (suffix == "txt" and content_type in {"text/plain", "application/octet-stream"})
        or (suffix == "pdf" and content_type == "application/pdf")
        or (
            suffix == "pptx"
            and content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )
    ):
        raise ApiError(
            415,
            "unsupported_source_file",
            "Upload a TXT, PDF, or PPTX file with its matching file type.",
        )
    try:
        if suffix == "txt":
            text = data.decode("utf-8-sig")
        elif suffix == "pdf":
            import io

            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        else:
            import io

            presentation = Presentation(io.BytesIO(data))
            text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
    except (UnicodeDecodeError, PdfReadError, PackageNotFoundError, ValueError):
        raise ApiError(422, "invalid_source_file", "This file could not be parsed.") from None
    text = text.strip()
    if not text:
        raise ApiError(422, "empty_source", "No readable text was found in this source file.")
    if len(text) > MAX_SOURCE_TEXT:
        text = text[:MAX_SOURCE_TEXT]
    return text


def _assert_public_host(hostname: str, port: int) -> None:
    """Resolve ``hostname`` and reject any private/local/mapped address.

    Called for the initial URL and again for every redirect target so a
    public URL cannot 302 to an internal host (SSRF via redirect).
    """
    try:
        addresses = socket.getaddrinfo(hostname, port)
    except socket.gaierror:
        raise ApiError(422, "invalid_source_url", "The source URL host could not be resolved.") from None
    if not addresses or any(
        not ipaddress.ip_address(address[4][0]).is_global for address in addresses
    ):
        raise ApiError(422, "invalid_source_url", "Private or local network URLs cannot be imported.")


def _fetch_url_text(url: str) -> str:
    max_redirects = 5
    headers = {
        "Accept": "text/html, text/plain",
        "User-Agent": SOURCE_FETCH_USER_AGENT,
    }
    try:
        with HttpClient(timeout=10, follow_redirects=False, trust_env=False) as http_client:
            current_url = url
            for _ in range(max_redirects + 1):
                parsed = urlsplit(current_url)
                if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
                    raise ApiError(422, "invalid_source_url", "Use a public HTTP or HTTPS URL.")
                _assert_public_host(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
                with http_client.stream("GET", current_url, headers=headers) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise ApiError(422, "source_fetch_failed", "The source URL could not be imported.")
                        current_url = urljoin(current_url, location)
                        continue
                    if response.status_code != 200:
                        raise ApiError(422, "source_fetch_failed", "The source URL could not be imported.")
                    media_type = response.headers.get("content-type", "").split(";", 1)[0].strip().lower()
                    chunks: list[bytes] = []
                    total = 0
                    for chunk in response.iter_bytes():
                        total += len(chunk)
                        if total > MAX_SOURCE_BYTES:
                            raise ApiError(413, "source_too_large", "The source URL content exceeds the 5 MB limit.")
                        chunks.append(chunk)
                    body = b"".join(chunks)
                    text = body.decode(response.encoding or "utf-8", errors="replace")
                    break
            else:
                raise ApiError(422, "source_fetch_failed", "The source URL could not be imported.")
        if media_type == "text/plain":
            pass
        elif media_type == "text/html":
            parser = TextExtractor()
            parser.feed(text)
            text = "".join(parser.parts)
        elif media_type == "application/pdf" or (
            media_type == "application/octet-stream"
            and urlsplit(current_url).path.lower().endswith(".pdf")
        ):
            text = _extract_file_text("source.pdf", "application/pdf", body)
        else:
            raise ApiError(
                415,
                "unsupported_source_url",
                "Only HTML pages, plain text, and PDF URLs can be imported.",
            )
        normalized = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        if not normalized:
            raise ApiError(422, "empty_source", "No readable text was found at this URL.")
        return normalized[:MAX_SOURCE_TEXT]
    except ApiError:
        raise
    except Exception:
        logger.exception("Unable to fetch instructor source URL.")
        raise ApiError(422, "source_fetch_failed", "The source URL could not be imported.") from None


@router.get("")
def list_materials(request: Request, client: OptionalClientDependency, store: StoreDependency):
    for user_id, db in _authorize_instructor(request, client, store):
        try:
            response = db.table("materials").select("*").eq("user_id", user_id).order(
                "updated_at", desc=True
            ).execute()
            rows = response.data or []
            ids = [row["id"] for row in rows]
            sources = []
            if ids:
                sources = db.table("material_sources").select("*").eq(
                    "user_id", user_id
                ).in_("material_id", ids).order("created_at").execute().data or []
            grouped: dict[str, list[dict[str, Any]]] = {}
            for source in sources:
                grouped.setdefault(str(source["material_id"]), []).append(source)
            return [
                _material_row(row, grouped.get(str(row["id"]), []), db)
                for row in rows
            ]
        except ApiError:
            raise
        except Exception as error:
            raise _database_error(error) from None


@router.post("", dependencies=[Depends(require_trusted_origin)])
def create_material(
    payload: MaterialCreate,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        now = datetime.now(UTC).isoformat()
        values = {
            "user_id": user_id,
            "title": payload.title,
            "type": payload.type,
            "data": payload.data,
            "chat_history": payload.chat_history,
            "updated_at": now,
        }
        if payload.id:
            values["id"] = str(payload.id)
        try:
            row = db.table("materials").insert(values).execute().data[0]
            return _material_row(row, [], db)
        except Exception as error:
            raise _database_error(error) from None


@router.patch("/{material_id}", dependencies=[Depends(require_trusted_origin)])
def update_material(
    material_id: uuid.UUID,
    payload: MaterialPatch,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        _owned_material(db, str(material_id), user_id)
        changes = payload.model_dump(by_alias=False, exclude_unset=True)
        if not changes:
            raise ApiError(422, "empty_update", "Provide a material field to update.")
        changes["updated_at"] = datetime.now(UTC).isoformat()
        try:
            row = db.table("materials").update(changes).eq(
                "id", str(material_id)
            ).eq("user_id", user_id).execute().data[0]
            sources = db.table("material_sources").select("*").eq(
                "material_id", str(material_id)
            ).eq("user_id", user_id).order("created_at").execute().data or []
            return _material_row(row, sources, db)
        except ApiError:
            raise
        except Exception as error:
            raise _database_error(error) from None


@router.delete("/{material_id}", status_code=204, dependencies=[Depends(require_trusted_origin)])
def delete_material(
    material_id: uuid.UUID,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        _owned_material(db, str(material_id), user_id)
        try:
            sources = db.table("material_sources").select("storage_path").eq(
                "material_id", str(material_id)
            ).eq("user_id", user_id).execute().data or []
            paths = [source["storage_path"] for source in sources if source.get("storage_path")]
            if paths:
                db.storage.from_(BUCKET_NAME).remove(paths)
            db.table("materials").delete().eq("id", str(material_id)).eq(
                "user_id", user_id
            ).execute()
        except Exception as error:
            raise _database_error(error) from None
        return None


class MaterialVersionRecord(BaseModel):
    id: str
    material_id: str = Field(alias="materialId")
    user_id: str = Field(alias="userId")
    title: str
    data: dict[str, Any]
    label: str
    created_at: str = Field(alias="createdAt")

    model_config = ConfigDict(populate_by_name=True)


class MaterialVersionCreate(BaseModel):
    label: str = Field(default="Snapshot", min_length=1, max_length=100)

    model_config = ConfigDict(populate_by_name=True)


def _version_row(row: dict[str, Any]) -> MaterialVersionRecord:
    return MaterialVersionRecord(
        id=str(row["id"]),
        materialId=str(row["material_id"]),
        userId=str(row["user_id"]),
        title=row["title"],
        data=row.get("data") or {},
        label=row.get("label") or "Snapshot",
        createdAt=row["created_at"],
    )


def _owned_version(db: Client, material_id: str, version_id: str, user_id: str) -> dict[str, Any]:
    try:
        result = db.table("material_versions").select("*").eq(
            "id", version_id
        ).eq("material_id", material_id).eq("user_id", user_id).maybe_single().execute()
    except Exception as error:
        raise _database_error(error) from None
    if not result or not result.data:
        raise ApiError(404, "version_not_found", "This saved version was not found.")
    return result.data


@router.post(
    "/{material_id}/versions",
    response_model=MaterialVersionRecord,
    dependencies=[Depends(require_trusted_origin)],
)
def create_version(
    material_id: uuid.UUID,
    payload: MaterialVersionCreate,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        material = _owned_material(db, str(material_id), user_id)
        try:
            row = db.table("material_versions").insert({
                "material_id": str(material_id),
                "user_id": user_id,
                "title": material["title"],
                "data": material.get("data") or {},
                "label": payload.label,
            }).execute().data[0]
            return _version_row(row)
        except Exception as error:
            raise _database_error(error) from None


@router.get("/{material_id}/versions", response_model=list[MaterialVersionRecord])
def list_versions(
    material_id: uuid.UUID,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        _owned_material(db, str(material_id), user_id)
        try:
            rows = db.table("material_versions").select("*").eq(
                "material_id", str(material_id)
            ).eq("user_id", user_id).order("created_at", desc=True).limit(25).execute().data or []
            return [_version_row(row) for row in rows]
        except Exception as error:
            raise _database_error(error) from None


@router.post(
    "/{material_id}/versions/{version_id}/restore",
    response_model=MaterialRecord,
    dependencies=[Depends(require_trusted_origin)],
)
def restore_version(
    material_id: uuid.UUID,
    version_id: uuid.UUID,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        material = _owned_material(db, str(material_id), user_id)
        version = _owned_version(db, str(material_id), str(version_id), user_id)
        try:
            # Snapshot the current state first so restoring a version is
            # itself reversible back to this exact moment.
            db.table("material_versions").insert({
                "material_id": str(material_id),
                "user_id": user_id,
                "title": material["title"],
                "data": material.get("data") or {},
                "label": "Before restore",
            }).execute()
            now = datetime.now(UTC).isoformat()
            row = db.table("materials").update({
                "data": version["data"],
                "updated_at": now,
            }).eq("id", str(material_id)).eq("user_id", user_id).execute().data[0]
            sources = db.table("material_sources").select("*").eq(
                "material_id", str(material_id)
            ).eq("user_id", user_id).order("created_at").execute().data or []
            return _material_row(row, sources, db)
        except ApiError:
            raise
        except Exception as error:
            raise _database_error(error) from None


@router.get("/{material_id}/export/pptx")
def export_material_pptx(
    material_id: uuid.UUID,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        material = _owned_material(db, str(material_id), user_id)
        try:
            file_bytes = build_pptx(material)
        except Exception:
            logger.exception("PPTX export failed for a material.")
            raise ApiError(500, "pptx_export_failed", "Unable to build the PPTX file. Try again.") from None
        title = re.sub(r'[\\/:*?"<>|]+', "", str(material.get("title") or "")).strip() or "material"
        return Response(
            content=file_bytes,
            media_type=PPTX_MEDIA_TYPE,
            headers={"Content-Disposition": f'attachment; filename="{title}.pptx"'},
        )


class MaterialPublish(BaseModel):
    published: bool


@router.post("/{material_id}/publish", response_model=MaterialRecord, dependencies=[Depends(require_trusted_origin)])
def publish_material(
    material_id: uuid.UUID,
    payload: MaterialPublish,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        _owned_material(db, str(material_id), user_id)
        now = datetime.now(UTC).isoformat()
        changes: dict[str, Any] = {
            "published": payload.published,
            "published_at": now if payload.published else None,
            "updated_at": now,
        }
        try:
            row = db.table("materials").update(changes).eq(
                "id", str(material_id)
            ).eq("user_id", user_id).execute().data[0]
            sources = db.table("material_sources").select("*").eq(
                "material_id", str(material_id)
            ).eq("user_id", user_id).order("created_at").execute().data or []
            return _material_row(row, sources, db)
        except Exception as error:
            raise _database_error(error) from None


@router.post(
    "/{material_id}/sources",
    response_model=SourceItem,
    dependencies=[Depends(require_trusted_origin)],
)
async def add_source(
    material_id: uuid.UUID,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        material = _owned_material(db, str(material_id), user_id)
        source_id = str(uuid.uuid4())
        storage_path: str | None = None
        file_data: bytes | None = None
        try:
            if request.headers.get("content-type", "").startswith("multipart/form-data"):
                form = await request.form()
                file = form.get("file")
                if not isinstance(file, StarletteUploadFile):
                    raise ApiError(422, "invalid_source", "Choose a source file to upload.")
                file_data = await file.read(MAX_SOURCE_BYTES + 1)
                if len(file_data) > MAX_SOURCE_BYTES:
                    raise ApiError(413, "source_too_large", "Source files must be 5 MB or smaller.")
                raw_content = _extract_file_text(file.filename or "source", file.content_type or "", file_data)
                name = (file.filename or "Uploaded source")[:255]
                suffix = name.rsplit(".", 1)[-1].lower()
                storage_path = f"{user_id}/{material_id}/{source_id}.{suffix}"
                source_type: SourceType = "file"
                source_url = None
                db.storage.from_(BUCKET_NAME).upload(
                    storage_path,
                    file_data,
                    file_options={"content-type": file.content_type or "application/octet-stream"},
                )
            else:
                payload = SourceCreate.model_validate(await request.json())
                name = payload.name
                source_type = payload.type
                source_url = str(payload.source_url) if payload.source_url else None
                if source_type == "url":
                    if not source_url:
                        raise ApiError(422, "invalid_source_url", "A source URL is required.")
                    raw_content = _fetch_url_text(source_url)
                else:
                    raw_content = (payload.raw_content or "").strip()
                    if not raw_content:
                        raise ApiError(422, "empty_source", "Enter text for this source.")
            row = db.table("material_sources").insert({
                "id": source_id,
                "material_id": str(material_id),
                "user_id": user_id,
                "type": source_type,
                "name": name,
                "raw_content": raw_content,
                "source_url": source_url,
                "storage_path": storage_path,
            }).execute().data[0]
            source = _material_row(
                material,
                [row],
                db,
            ).sources[0]
            return source
        except ApiError:
            raise
        except Exception as error:
            if storage_path:
                try:
                    db.storage.from_(BUCKET_NAME).remove([storage_path])
                except Exception:
                    logger.exception("Unable to clean up an incomplete source upload.")
            raise _database_error(error) from None


@router.delete(
    "/{material_id}/sources/{source_id}",
    status_code=204,
    dependencies=[Depends(require_trusted_origin)],
)
def delete_source(
    material_id: uuid.UUID,
    source_id: uuid.UUID,
    request: Request,
    client: OptionalClientDependency,
    store: StoreDependency,
):
    for user_id, db in _authorize_instructor(request, client, store):
        _owned_material(db, str(material_id), user_id)
        try:
            result = db.table("material_sources").select("*").eq(
                "id", str(source_id)
            ).eq("material_id", str(material_id)).eq("user_id", user_id).maybe_single().execute()
            if not result or not result.data:
                raise ApiError(404, "source_not_found", "This material source was not found.")
            storage_path = result.data.get("storage_path")
            if storage_path:
                db.storage.from_(BUCKET_NAME).remove([storage_path])
            db.table("material_sources").delete().eq("id", str(source_id)).eq(
                "material_id", str(material_id)
            ).eq("user_id", user_id).execute()
        except ApiError:
            raise
        except Exception as error:
            raise _database_error(error) from None
        return None
