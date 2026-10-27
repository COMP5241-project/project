"""Build PowerPoint files from instructor materials.

Used by the PPTX export endpoint. The layout is intentionally simple —
title, bullets, the visual (steps or photo), on 16:9 slides using the
material's theme colors — and mirrors the on-screen slide canvas closely
enough that an exported deck looks like the instructor's deck.
"""

import base64
import re
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

PPTX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

_DEFAULT_THEME = {
    "background": "#FFFFFF",
    "textColor": "#0F172A",
    "accentColor": "#A21942",
    "cardBackground": "#F8FAFC",
}

_HEX_COLOR = re.compile(r"^#?([0-9a-fA-F]{6})$")


def _rgb(value: str | None, fallback: str) -> RGBColor:
    match = _HEX_COLOR.match((value or "").strip())
    if not match:
        match = _HEX_COLOR.match(fallback)
    return RGBColor.from_string(match.group(1).upper())


def _theme_for(slide_data: dict[str, Any]) -> dict[str, str]:
    raw = slide_data.get("theme") or {}
    return {
        "background": str(raw.get("background") or _DEFAULT_THEME["background"]),
        "textColor": str(raw.get("textColor") or _DEFAULT_THEME["textColor"]),
        "accentColor": str(raw.get("accentColor") or _DEFAULT_THEME["accentColor"]),
        "cardBackground": str(raw.get("cardBackground") or _DEFAULT_THEME["cardBackground"]),
    }


def _add_textbox(slide: Any, left: float, top: float, width: float, height: float) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    return box.text_frame


def _set_background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color, _DEFAULT_THEME["background"])


def _build_slide(prs: Presentation, slide_data: dict[str, Any]) -> None:
    theme = _theme_for(slide_data)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, theme["background"])

    title = str(slide_data.get("title") or "Untitled slide")
    title_frame = _add_textbox(slide, 0.7, 0.45, 11.9, 1.6)
    title_run = title_frame.paragraphs[0].add_run()
    title_run.text = title
    title_run.font.size = Pt(40)
    title_run.font.bold = True
    title_run.font.color.rgb = _rgb(theme["textColor"], _DEFAULT_THEME["textColor"])

    bullets = [str(bullet) for bullet in (slide_data.get("bullets") or []) if str(bullet).strip()]
    if bullets:
        bullet_frame = _add_textbox(slide, 0.7, 2.4, 6.0, 4.3)
        for index, bullet in enumerate(bullets):
            paragraph = bullet_frame.paragraphs[0] if index == 0 else bullet_frame.add_paragraph()
            run = paragraph.add_run()
            run.text = f"• {bullet}"
            run.font.size = Pt(20)
            run.font.color.rgb = _rgb(theme["textColor"], _DEFAULT_THEME["textColor"])
            paragraph.space_after = Pt(10)

    visual = slide_data.get("visual") or {}
    visual_left = 7.3
    visual_top = 2.4
    image_data = slide_data.get("imageData")
    photo_embedded = False
    if visual.get("kind") == "photo" and isinstance(image_data, str) and image_data.startswith("data:"):
        try:
            encoded = image_data.split(",", 1)[1]
            picture = slide.shapes.add_picture(
                BytesIO(base64.b64decode(encoded)),
                Inches(visual_left),
                Inches(visual_top),
                height=Inches(4.3),
            )
            max_width = Inches(5.3)
            if picture.width > max_width:
                scale = max_width / picture.width
                picture.width = max_width
                picture.height = int(picture.height * scale)
            photo_embedded = True
        except Exception:
            photo_embedded = False  # fall back to the steps card

    if not photo_embedded:
        steps = [str(step) for step in (visual.get("steps") or []) if str(step).strip()]
        if steps:
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(visual_left),
                Inches(visual_top),
                Inches(5.3),
                Inches(4.3),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = _rgb(theme["cardBackground"], _DEFAULT_THEME["cardBackground"])
            card.line.color.rgb = _rgb(theme["accentColor"], _DEFAULT_THEME["accentColor"])
            frame = card.text_frame
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for index, step in enumerate(steps[:4]):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                run = paragraph.add_run()
                run.text = f"{index + 1}. {step}"
                run.font.size = Pt(16)
                run.font.bold = index == 0
                run.font.color.rgb = _rgb(theme["textColor"], _DEFAULT_THEME["textColor"])
                paragraph.space_after = Pt(8)


def _build_quiz(prs: Presentation, material: dict[str, Any], data: dict[str, Any]) -> None:
    heading = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(heading, _DEFAULT_THEME["background"])
    heading_frame = _add_textbox(heading, 0.7, 2.7, 11.9, 2.2)
    heading_run = heading_frame.paragraphs[0].add_run()
    heading_run.text = str(material.get("title") or "Quiz")
    heading_run.font.size = Pt(40)
    heading_run.font.bold = True
    heading_run.font.color.rgb = _rgb(_DEFAULT_THEME["textColor"], _DEFAULT_THEME["textColor"])

    for index, question in enumerate(data.get("questions") or []):
        if not isinstance(question, dict):
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(slide, _DEFAULT_THEME["background"])

        label_frame = _add_textbox(slide, 0.7, 0.5, 11.9, 0.7)
        label_run = label_frame.paragraphs[0].add_run()
        label_run.text = f"Question {index + 1}"
        label_run.font.size = Pt(20)
        label_run.font.bold = True
        label_run.font.color.rgb = _rgb(_DEFAULT_THEME["accentColor"], _DEFAULT_THEME["accentColor"])

        question_frame = _add_textbox(slide, 0.7, 1.2, 11.9, 1.4)
        question_run = question_frame.paragraphs[0].add_run()
        question_run.text = str(question.get("question") or "")
        question_run.font.size = Pt(26)
        question_run.font.color.rgb = _rgb(_DEFAULT_THEME["textColor"], _DEFAULT_THEME["textColor"])

        options = [str(option) for option in (question.get("options") or []) if str(option).strip()]
        if options:
            options_frame = _add_textbox(slide, 0.9, 2.9, 11.4, 4.0)
            for option_index, option in enumerate(options):
                paragraph = options_frame.paragraphs[0] if option_index == 0 else options_frame.add_paragraph()
                run = paragraph.add_run()
                run.text = f"{chr(65 + option_index)}. {option}"
                run.font.size = Pt(18)
                run.font.color.rgb = _rgb(_DEFAULT_THEME["textColor"], _DEFAULT_THEME["textColor"])
                paragraph.space_after = Pt(10)

        # The answer key stays in speaker notes, never on the slide itself.
        correct_index = question.get("correctAnswerIndex")
        explanation = question.get("explanation") or ""
        notes_parts = []
        if isinstance(correct_index, int) and 0 <= correct_index < len(options):
            notes_parts.append(f"Correct answer: {chr(65 + correct_index)}.")
        if explanation:
            notes_parts.append(str(explanation))
        if notes_parts:
            slide.notes_slide.notes_text_frame.text = " ".join(notes_parts)


def build_pptx(material: dict[str, Any]) -> bytes:
    """Builds a PPTX file from a material row (see the ``materials`` table)."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)

    data = material.get("data") or {}
    if data.get("type") == "quiz":
        _build_quiz(prs, material, data)
    else:
        for slide_data in data.get("content") or []:
            if isinstance(slide_data, dict):
                _build_slide(prs, slide_data)

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
