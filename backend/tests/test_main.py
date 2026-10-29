from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from threading import Event
from types import SimpleNamespace
import time
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from httpx import Client as HttpxClient
from postgrest.exceptions import APIError as PostgrestAPIError
from pptx import Presentation
from storage3.exceptions import StorageApiError

from app.database.supabase import get_supabase
from app.identity.errors import ApiError
from app.course_generation import (
    CourseGeneratorNotConfiguredError,
    CourseGeneratorUnavailableError,
    GeneratedCourseContent,
    GeneratedQuizContent,
    generate_course_content,
)
from app.routers.materials import _extract_file_text, _fetch_url_text
from app.pptx_export import build_pptx
from app.hallucination.judge import JudgeUnavailableError, JudgeVerdict, parse_verdict
from app.identity.auth import get_auth_client, get_optional_auth_client
from app.identity.sessions import COOKIE_NAME, Session, SessionStore, get_session_store
from app.main import app
from app.routers.hallucination import judge


@pytest.fixture
def mock_supabase() -> Mock:
    return Mock()


@pytest.fixture
def auth_client() -> Mock:
    client = Mock()
    client.options.headers = {}
    user = SimpleNamespace(id="student-id", email="student@example.com")
    session = SimpleNamespace(
        access_token="private-access-token", refresh_token="private-refresh-token",
        expires_at=time.time() + 3600, expires_in=3600,
    )
    client.auth.sign_in_with_password.return_value = SimpleNamespace(user=user, session=session)
    client.auth.refresh_session.return_value = SimpleNamespace(user=user, session=session)
    client.auth.get_user.return_value = SimpleNamespace(user=user)
    client.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": user.id, "email": user.email, "role": "student",
        "cs_qualified": True, "assessment_completed": False,
    }
    return client


@pytest.fixture
def store() -> SessionStore:
    return SessionStore()


@pytest.fixture
def client(mock_supabase: Mock, auth_client: Mock, store: SessionStore) -> Iterator[HttpxClient]:
    app.dependency_overrides[get_supabase] = lambda: mock_supabase
    app.dependency_overrides[get_auth_client] = lambda: auth_client
    app.dependency_overrides[get_optional_auth_client] = lambda: auth_client
    app.dependency_overrides[get_session_store] = lambda: store
    with TestClient(app) as test_client:
        test_client.headers["Origin"] = "http://localhost:5173"
        yield test_client
    app.dependency_overrides.clear()


def test_root_returns_greeting(client: HttpxClient) -> None:
    response = client.get("/")

    assert response.status_code == 200
    assert response.json() == {"message": "Hello"}


def test_health_returns_ok_when_supabase_query_succeeds(
    client: HttpxClient, mock_supabase: Mock
) -> None:
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok"}
    mock_supabase.table.assert_called_once_with("profiles")
    mock_supabase.table.return_value.select.assert_called_once_with("*")
    mock_supabase.table.return_value.select.return_value.limit.assert_called_once_with(1)
    mock_supabase.table.return_value.select.return_value.limit.return_value.execute.assert_called_once()


def test_health_returns_service_unavailable_when_supabase_query_fails(
    client: HttpxClient, mock_supabase: Mock
) -> None:
    execute = mock_supabase.table.return_value.select.return_value.limit.return_value.execute
    execute.side_effect = RuntimeError("database unavailable")

    response = client.get("/health")

    assert response.status_code == 503
    assert response.json() == {
        "status": "error",
        "message": "database unavailable",
    }


def auth_login(client: HttpxClient):
    return client.post("/api/auth/login", json={"email": "student@example.com", "password": "secret-password"})


def test_auth_login_and_profile_use_private_cookie(client: HttpxClient, auth_client: Mock):
    response = auth_login(client)
    assert response.status_code == 200
    assert response.json()["profile"]["role"] == "student"
    assert "private-" not in response.text
    assert "HttpOnly" in response.headers["set-cookie"]
    assert "SameSite=lax" in response.headers["set-cookie"]
    assert "Path=/api" in response.headers["set-cookie"]
    assert response.headers["cache-control"] == "no-store"
    assert client.cookies.get(COOKIE_NAME) != "private-access-token"
    response = client.get("/api/profiles/me")
    assert response.status_code == 200
    auth_client.auth.get_user.assert_called_once_with("private-access-token")
    auth_client.postgrest.auth.assert_called_with("private-access-token")
    auth_client.table.return_value.select.return_value.eq.assert_called_with("id", "student-id")


def test_password_reset_request_sends_email_otp(client: HttpxClient, auth_client: Mock):
    response = client.post(
        "/api/auth/password-reset/request",
        json={"email": "student@example.com"},
    )

    assert response.status_code == 204
    auth_client.auth.sign_in_with_otp.assert_called_once_with({
        "email": "student@example.com",
        "options": {"should_create_user": False},
    })


def test_password_reset_verify_updates_password(client: HttpxClient, auth_client: Mock):
    response = client.post(
        "/api/auth/password-reset/verify",
        json={
            "email": "student@example.com",
            "code": "12345678",
            "password": "new-password",
            "password_confirmation": "new-password",
        },
    )

    assert response.status_code == 204
    auth_client.auth.verify_otp.assert_called_once_with({
        "email": "student@example.com",
        "token": "12345678",
        "type": "email",
    })
    auth_client.auth.update_user.assert_called_once_with({"password": "new-password"})


def test_password_reset_rejects_mismatched_passwords(client: HttpxClient, auth_client: Mock):
    response = client.post(
        "/api/auth/password-reset/verify",
        json={
            "email": "student@example.com",
            "code": "12345678",
            "password": "new-password",
            "password_confirmation": "different-password",
        },
    )

    assert response.status_code == 422
    auth_client.auth.verify_otp.assert_not_called()
    auth_client.auth.update_user.assert_not_called()


def test_password_reset_rejects_six_digit_codes(client: HttpxClient, auth_client: Mock):
    response = client.post(
        "/api/auth/password-reset/verify",
        json={
            "email": "student@example.com",
            "code": "123456",
            "password": "new-password",
            "password_confirmation": "new-password",
        },
    )

    assert response.status_code == 422
    auth_client.auth.verify_otp.assert_not_called()
    auth_client.auth.update_user.assert_not_called()


def test_password_reset_does_not_update_when_code_verification_fails(
    client: HttpxClient, auth_client: Mock
):
    auth_client.auth.verify_otp.side_effect = RuntimeError("invalid code")

    response = client.post(
        "/api/auth/password-reset/verify",
        json={
            "email": "student@example.com",
            "code": "12345678",
            "password": "new-password",
            "password_confirmation": "new-password",
        },
    )

    assert response.status_code == 503
    auth_client.auth.update_user.assert_not_called()



def test_auth_missing_session(client: HttpxClient):
    response = client.get("/api/profiles/me")
    assert response.status_code == 401
    assert response.json()["error"]["code"] == "unauthenticated"


@pytest.mark.parametrize("origin", ["https://evil.example", ""])
def test_auth_rejects_untrusted_origin(client: HttpxClient, auth_client: Mock, origin: str):
    response = client.post("/api/auth/login", json={"email": "x", "password": "secret"}, headers={"Origin": origin})
    assert response.status_code == 403
    auth_client.auth.sign_in_with_password.assert_not_called()


@pytest.mark.parametrize(
    "origin",
    ["http://localhost:5173", "http://127.0.0.1:5173", "http://[::1]:5173", "http://LOCALHOST:5173"],
)
def test_auth_accepts_all_loopback_origin_spellings(
    client: HttpxClient, monkeypatch: pytest.MonkeyPatch, origin: str
):
    monkeypatch.setenv("FRONTEND_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173")
    response = client.post(
        "/api/auth/login",
        json={"email": "student@example.com", "password": "secret-password"},
        headers={"Origin": origin},
    )
    assert response.status_code == 200


def test_auth_rejects_wrong_port_origin(
    client: HttpxClient, auth_client: Mock, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("FRONTEND_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173")
    response = client.post(
        "/api/auth/login", json={"email": "x", "password": "secret"},
        headers={"Origin": "http://localhost:5174"},
    )
    assert response.status_code == 403
    auth_client.auth.sign_in_with_password.assert_not_called()


def test_auth_accepts_current_codespaces_frontend_origin(
    client: HttpxClient, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("FRONTEND_ORIGINS", "http://localhost:5173")
    monkeypatch.setenv("CODESPACE_NAME", "example-codespace")
    monkeypatch.setenv("GITHUB_CODESPACES_PORT_FORWARDING_DOMAIN", "app.github.dev")
    response = client.post(
        "/api/auth/login",
        json={"email": "student@example.com", "password": "secret-password"},
        headers={"Origin": "https://example-codespace-5173.app.github.dev"},
    )
    assert response.status_code == 200


def test_auth_accepts_any_forwarded_codespaces_frontend_port(
    client: HttpxClient, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("FRONTEND_ORIGINS", "http://localhost:5173")
    monkeypatch.setenv("CODESPACE_NAME", "example-codespace")
    monkeypatch.setenv("GITHUB_CODESPACES_PORT_FORWARDING_DOMAIN", "app.github.dev")
    response = client.post(
        "/api/auth/login",
        json={"email": "student@example.com", "password": "secret-password"},
        headers={"Origin": "https://example-codespace-5180.app.github.dev"},
    )
    assert response.status_code == 200


def test_auth_accepts_explicit_https_default_port_for_codespaces_origin(
    client: HttpxClient, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("FRONTEND_ORIGINS", "http://localhost:5173")
    monkeypatch.setenv("CODESPACE_NAME", "example-codespace")
    monkeypatch.setenv("GITHUB_CODESPACES_PORT_FORWARDING_DOMAIN", "app.github.dev")
    response = client.post(
        "/api/auth/login",
        json={"email": "student@example.com", "password": "secret-password"},
        headers={"Origin": "https://example-codespace-5173.app.github.dev:443"},
    )
    assert response.status_code == 200


def test_auth_rejects_other_codespaces_frontend_origin(
    client: HttpxClient, auth_client: Mock, monkeypatch: pytest.MonkeyPatch
):
    monkeypatch.setenv("FRONTEND_ORIGINS", "http://localhost:5173")
    monkeypatch.setenv("CODESPACE_NAME", "example-codespace")
    monkeypatch.setenv("GITHUB_CODESPACES_PORT_FORWARDING_DOMAIN", "app.github.dev")
    response = client.post(
        "/api/auth/login",
        json={"email": "x", "password": "secret"},
        headers={"Origin": "https://other-codespace-5173.app.github.dev"},
    )
    assert response.status_code == 403
    auth_client.auth.sign_in_with_password.assert_not_called()


def test_auth_invalid_request_does_not_echo_password(client: HttpxClient):
    response = client.post("/api/auth/login", json={"email": "", "password": "secret-password"})
    assert response.status_code == 422
    assert "secret-password" not in response.text


def test_auth_invalid_credentials(client: HttpxClient, auth_client: Mock):
    from supabase_auth.errors import AuthApiError
    auth_client.auth.sign_in_with_password.side_effect = AuthApiError("provider details", 400, "invalid_credentials")
    response = auth_login(client)
    assert response.status_code == 401
    assert "provider details" not in response.text
    assert COOKIE_NAME not in client.cookies


@pytest.mark.parametrize("data,status", [(None, 404), ({"role": "admin"}, 403), ("malformed", 503)])
def test_auth_profile_failure_does_not_create_session(client: HttpxClient, auth_client: Mock, store: SessionStore, data, status):
    auth_client.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = data
    response = auth_login(client)
    assert response.status_code == status
    assert not store.sessions
    assert COOKIE_NAME not in client.cookies
    auth_client.auth.admin.sign_out.assert_called_once_with("private-access-token", "local")


@pytest.mark.parametrize("data,status,code", [
    (None, 404, "profile_missing"),
    ({"role": "admin"}, 403, "unsupported_role"),
    ("malformed", 503, "profile_unavailable"),
])
def test_auth_profile_failure_preserves_session(client: HttpxClient, auth_client: Mock, store: SessionStore, data, status, code):
    auth_login(client)
    session_id = client.cookies.get(COOKIE_NAME)
    session = store.get(session_id)
    auth_client.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = data
    response = client.get("/api/profiles/me")
    assert response.status_code == status
    assert response.json()["error"]["code"] == code
    assert store.get(session_id) is session
    assert client.cookies.get(COOKIE_NAME) == session_id
    auth_client.auth.admin.sign_out.assert_not_called()


def test_auth_invalidated_session_during_profile_load_is_rejected(client: HttpxClient, auth_client: Mock, store: SessionStore):
    auth_login(client)
    session_id = client.cookies.get(COOKIE_NAME)
    execute = auth_client.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute
    result = execute.return_value

    def invalidate_during_profile_load():
        store.remove_user("student-id")
        return result

    execute.side_effect = invalidate_during_profile_load
    response = client.get("/api/profiles/me")
    assert response.status_code == 401
    assert response.json()["error"]["code"] == "unauthenticated"
    assert store.get(session_id) is None
    assert COOKIE_NAME not in client.cookies


def test_auth_refresh_rotates_tokens(client: HttpxClient, auth_client: Mock, store: SessionStore):
    auth_login(client)
    session = store.get(client.cookies.get(COOKIE_NAME))
    session.token_expires_at = 0
    refreshed = auth_client.auth.refresh_session.return_value.session
    refreshed.access_token = "rotated-access"
    refreshed.refresh_token = "rotated-refresh"
    response = client.get("/api/profiles/me")
    assert response.status_code == 200
    auth_client.auth.refresh_session.assert_called_once_with("private-refresh-token")
    auth_client.auth.get_user.assert_called_with("rotated-access")
    assert session.refresh_token == "rotated-refresh"


@pytest.mark.parametrize("invalid", [True, False])
def test_auth_refresh_error_distinguishes_expiry_from_outage(client: HttpxClient, auth_client: Mock, store: SessionStore, invalid: bool):
    from supabase_auth.errors import AuthApiError
    auth_login(client)
    session_id = client.cookies.get(COOKIE_NAME)
    store.get(session_id).token_expires_at = 0
    auth_client.auth.refresh_session.side_effect = AuthApiError("expired", 400, "refresh_token_not_found") if invalid else RuntimeError("upstream offline")
    response = client.get("/api/profiles/me")
    assert response.status_code == (401 if invalid else 503)
    assert (store.get(session_id) is None) == invalid
    assert "upstream offline" not in response.text


def test_auth_absolute_session_expiry(client: HttpxClient, store: SessionStore):
    auth_login(client)
    store.get(client.cookies.get(COOKIE_NAME)).expires_at = 0
    assert client.get("/api/profiles/me").status_code == 401


def test_auth_replaced_session_is_removed(client: HttpxClient, store: SessionStore):
    auth_login(client)
    old_id = client.cookies.get(COOKIE_NAME)
    auth_login(client)
    assert store.get(old_id) is None
    assert len(store.sessions) == 1


def test_auth_rejects_mismatched_user(client: HttpxClient, auth_client: Mock, store: SessionStore):
    auth_login(client)
    auth_client.auth.get_user.return_value.user = SimpleNamespace(id="other-user", email="other@example.com")
    assert client.get("/api/profiles/me").status_code == 401
    assert not store.sessions


def test_auth_logout_without_session(client: HttpxClient):
    assert client.post("/api/auth/logout", json={}).status_code == 204


@pytest.mark.parametrize("unavailable", [False, True])
def test_auth_logout_invalidates_sessions_even_when_revocation_fails(client: HttpxClient, auth_client: Mock, store: SessionStore, unavailable: bool):
    auth_login(client)
    if unavailable:
        auth_client.auth.admin.sign_out.side_effect = RuntimeError("provider-secret")
    response = client.post("/api/auth/logout", json={})
    assert response.status_code == (503 if unavailable else 204)
    if unavailable:
        assert response.json()["error"]["signed_out"] is True
        assert "provider-secret" not in response.text
    assert not store.sessions
    assert COOKIE_NAME not in client.cookies
    assert client.get("/api/profiles/me").status_code == 401
    auth_client.auth.admin.sign_out.assert_called_once_with("private-access-token", "global")


def test_auth_logout_clears_only_this_users_sessions(client: HttpxClient, store: SessionStore):
    auth_login(client)
    store.add(Session("student-id", "second-token", "second-refresh", time.time() + 3600, time.time() + 3600))
    other_id = store.add(Session("other-user", "other-token", "other-refresh", time.time() + 3600, time.time() + 3600))
    assert client.post("/api/auth/logout", json={}).status_code == 204
    assert len(store.sessions) == 1
    assert store.get(other_id).user_id == "other-user"


def test_auth_refresh_is_serialized(client: HttpxClient, auth_client: Mock, store: SessionStore):
    auth_login(client)
    session_id = client.cookies.get(COOKIE_NAME)
    store.get(session_id).token_expires_at = 0
    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(executor.map(lambda _: client.get("/api/profiles/me"), range(2)))
    assert all(response.status_code == 200 for response in responses)
    auth_client.auth.refresh_session.assert_called_once()


def test_auth_logout_does_not_resurrect_inflight_session(client: HttpxClient, auth_client: Mock, store: SessionStore):
    auth_login(client)
    session_id = client.cookies.get(COOKIE_NAME)
    started = Event()
    release = Event()
    verified = auth_client.auth.get_user.return_value

    def delayed_user(token):
        started.set()
        assert release.wait(5)
        return verified

    auth_client.auth.get_user.side_effect = delayed_user
    with ThreadPoolExecutor(max_workers=1) as executor:
        pending = executor.submit(client.get, "/api/profiles/me")
        assert started.wait(5)
        store.remove_user("student-id")
        release.set()
        assert pending.result(timeout=5).status_code == 401
    assert store.get(session_id) is None


def test_auth_missing_configuration_does_not_block_local_logout(client: HttpxClient, store: SessionStore):
    auth_login(client)
    app.dependency_overrides[get_optional_auth_client] = lambda: None
    response = client.post("/api/auth/logout", json={})
    assert response.status_code == 503
    assert response.json()["error"]["signed_out"] is True
    assert not store.sessions
    assert client.get("/api/profiles/me").status_code == 401


def test_auth_secure_cookie_configuration(client: HttpxClient, monkeypatch):
    monkeypatch.setenv("SESSION_COOKIE_SECURE", "true")
    assert "Secure" in auth_login(client).headers["set-cookie"]


def test_auth_mutations_require_json(client: HttpxClient):
    assert client.post("/api/auth/logout", content="text").status_code == 415


def test_hallucination_result_is_saved_for_signed_in_user(client: HttpxClient, auth_client: Mock):
    auth_login(client)
    response = client.post("/api/profiles/hallucination-results", json={
        "challenge_id": "off-by-one", "score": 80, "critique": "The loop stops one short.",
        "feedback": "Good eye!",
    })
    assert response.status_code == 204
    auth_client.postgrest.auth.assert_called_with("private-access-token")
    auth_client.table.return_value.insert.assert_called_once_with({
        "user_id": "student-id", "challenge_id": "off-by-one", "score": 80,
        "critique": "The loop stops one short.", "feedback": "Good eye!",
    })


def test_hallucination_result_requires_session(client: HttpxClient, auth_client: Mock):
    response = client.post("/api/profiles/hallucination-results", json={
        "challenge_id": "off-by-one", "score": 80, "critique": "A long enough critique.",
    })
    assert response.status_code == 401
    auth_client.table.return_value.insert.assert_not_called()


def test_hallucination_result_validation(client: HttpxClient):
    auth_login(client)
    response = client.post("/api/profiles/hallucination-results", json={
        "challenge_id": "off-by-one", "score": 101, "critique": "A long enough critique.",
    })
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_request"


def test_hallucination_result_save_failure_is_service_unavailable(client: HttpxClient, auth_client: Mock):
    auth_login(client)
    auth_client.table.return_value.insert.return_value.execute.side_effect = RuntimeError("database unavailable")
    response = client.post("/api/profiles/hallucination-results", json={
        "challenge_id": "off-by-one", "score": 80, "critique": "A long enough critique.",
    })
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "result_save_failed"


def test_auth_clients_do_not_share_authentication_state(monkeypatch):
    from httpx import Client as HttpClient
    from app.database.supabase import create_user_client
    monkeypatch.setenv("SUPABASE_URL", "https://example.supabase.co")
    monkeypatch.setenv("SUPABASE_KEY", "public-test-key")
    with HttpClient() as first_transport, HttpClient() as second_transport:
        first = create_user_client(first_transport)
        second = create_user_client(second_transport)
        first.postgrest.auth("first-user-jwt")
        second.postgrest.auth("second-user-jwt")
        assert first.postgrest.headers["Authorization"] == "Bearer first-user-jwt"
        assert second.postgrest.headers["Authorization"] == "Bearer second-user-jwt"
        assert first.options.persist_session is False
        assert first.options.auto_refresh_token is False


JUDGE_PAYLOAD = {
    "challenge_id": "ncs-01", "critique": "This function does not exist.", "tapped_line": 4,
}


def test_hallucination_challenges_hide_answers(client: HttpxClient):
    response = client.get("/api/hallucination/challenges")
    assert response.status_code == 200
    challenges = response.json()
    assert [item["id"] for item in challenges] == ["ncs-01", "ncs-02", "ncs-03", "ncs-04"]
    for item in challenges:
        assert "flaw_line" not in item
        assert "flaw_summary" not in item
        assert item["code"] and item["hint_mild"] and item["hint_strong"]


def test_hallucination_judge_requires_session(client: HttpxClient, auth_client: Mock):
    response = client.post("/api/hallucination/judge", json=JUDGE_PAYLOAD)
    assert response.status_code == 401
    auth_client.auth.get_user.assert_not_called()


def test_hallucination_judge_rejects_unknown_challenge(client: HttpxClient, monkeypatch):
    auth_login(client)
    monkeypatch.setattr(
        "app.routers.hallucination.judge_critique",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("must not be called")),
    )
    response = client.post(
        "/api/hallucination/judge",
        json={"challenge_id": "ncs-99", "critique": "A long enough critique.", "tapped_line": 0},
    )
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "challenge_not_found"


def test_hallucination_judge_returns_verdict_with_reveal(client: HttpxClient, monkeypatch):
    auth_login(client)
    verdict = JudgeVerdict(
        score=90, rating="sharp", identified=["missing function"],
        missed=[], feedback="Great detective work.",
    )
    monkeypatch.setattr(
        "app.routers.hallucination.judge_critique",
        lambda *args, **kwargs: verdict,
    )
    response = client.post("/api/hallucination/judge", json=JUDGE_PAYLOAD)
    assert response.status_code == 200
    body = response.json()
    assert body["score"] == 90 and body["rating"] == "sharp"
    assert body["flaw_line"] == 4
    assert body["flaw_summary"].startswith("calendar.set_month_name()")


def test_hallucination_judge_upstream_failure_is_service_unavailable(client: HttpxClient, monkeypatch):
    auth_login(client)

    def unavailable(*args, **kwargs):
        raise JudgeUnavailableError("boom")

    monkeypatch.setattr("app.routers.hallucination.judge_critique", unavailable)
    response = client.post("/api/hallucination/judge", json=JUDGE_PAYLOAD)
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "judge_unavailable"


def test_instructor_course_generation_requires_session(client: HttpxClient):
    response = client.post(
        "/api/instructor/generate",
        json={"prompt": "Create an introduction to Python."},
    )
    assert response.status_code == 401


def test_instructor_course_generation_rejects_students(
    client: HttpxClient, auth_client: Mock, monkeypatch
):
    monkeypatch.setattr(
        "app.routers.instructor.generate_course_content",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("must not be called")),
    )
    auth_login(client)
    response = client.post(
        "/api/instructor/generate",
        json={"prompt": "Create an introduction to Python."},
    )
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "instructor_required"


def test_instructor_course_generation_returns_slides(
    client: HttpxClient, auth_client: Mock, monkeypatch
):
    auth_client.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    auth_login(client)
    generated = GeneratedCourseContent.model_validate({
        "type": "slides",
        "content": [{
            "title": "Python Basics",
            "bullets": ["Variables", "Functions"],
            "visual": {
                "kind": "diagram",
                "description": "From variables to functions",
                "steps": ["Variables", "Functions"],
            },
        }],
    })
    generate = Mock(return_value=generated)
    monkeypatch.setattr("app.routers.instructor.generate_course_content", generate)

    response = client.post(
        "/api/instructor/generate",
        json={
            "prompt": "Create an introduction to Python.",
            "source_context": "Course notes about variables and functions.",
        },
    )

    assert response.status_code == 200
    assert response.json() == generated.model_dump()
    generate.assert_called_once()
    assert generate.call_args.args[:2] == (
        "Create an introduction to Python.",
        "Course notes about variables and functions.",
    )
    assert generate.call_args.kwargs == {
        "artifact_type": "slides",
        "existing_content": None,
        "revision_prompt": "",
    }


def test_instructor_course_generation_returns_quiz(
    client: HttpxClient, auth_client: Mock, monkeypatch
):
    auth_client.table.return_value.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    auth_login(client)
    generated = GeneratedQuizContent.model_validate({
        "type": "quiz",
        "questions": [{
            "id": "q1",
            "question": "Which type stores a whole number?",
            "options": ["int", "str"],
            "correctAnswerIndex": 0,
            "explanation": "An int stores an integer value.",
        }],
    })
    generate = Mock(return_value=generated)
    monkeypatch.setattr("app.routers.instructor.generate_course_content", generate)

    response = client.post(
        "/api/instructor/generate",
        json={"prompt": "Create a Python multiple choice quiz."},
    )

    assert response.status_code == 200
    assert response.json() == generated.model_dump(by_alias=True)
    assert generate.call_args.kwargs == {
        "artifact_type": "quiz",
        "existing_content": None,
        "revision_prompt": "",
    }


def test_instructor_course_generation_validates_request(client: HttpxClient):
    response = client.post("/api/instructor/generate", json={"prompt": ""})
    assert response.status_code == 422


def test_materials_api_lists_account_materials_with_nested_sources(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    source_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
        "material_sources": source_table,
    }[name]
    auth_login(client)
    material_table.select.return_value.eq.return_value.order.return_value.execute.return_value.data = [{
        "id": "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817",
        "user_id": "student-id",
        "title": "Python deck",
        "type": "slides",
        "data": {"type": "slides", "content": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }]
    source_table.select.return_value.eq.return_value.in_.return_value.order.return_value.execute.return_value.data = [{
        "id": "9be17748-cdad-4b87-84f2-032a8a23a11b",
        "material_id": "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817",
        "user_id": "student-id",
        "type": "text",
        "name": "Lecture notes",
        "raw_content": "Functions and variables.",
        "source_url": None,
        "storage_path": None,
        "created_at": "2026-10-15T12:00:00+00:00",
    }]

    response = client.get("/api/materials")

    assert response.status_code == 200
    material = response.json()[0]
    assert material["userId"] == "student-id"
    assert material["sources"][0]["materialId"] == material["id"]
    assert material["sources"][0]["rawContent"] == "Functions and variables."
    material_table.select.return_value.eq.assert_called_once_with("user_id", "student-id")
    source_table.select.return_value.eq.assert_called_once_with("user_id", "student-id")
    assert auth_client.options.headers["Authorization"] == "Bearer private-access-token"


def test_materials_api_explains_when_database_migration_is_missing(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
    }[name]
    auth_login(client)
    material_table.select.return_value.eq.return_value.order.return_value.execute.side_effect = (
        PostgrestAPIError({
            "message": "Could not find the table 'public.materials' in the schema cache",
            "code": "PGRST205",
        })
    )

    response = client.get("/api/materials")

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "materials_not_configured"
    assert "20261015000000_create_instructor_materials.sql" in response.json()["error"]["message"]


def test_material_source_text_is_saved_only_to_its_owned_material(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    source_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
        "material_sources": source_table,
    }[name]
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Python deck",
        "type": "slides",
        "data": {"type": "slides", "content": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    source_table.insert.return_value.execute.return_value.data = [{
        "id": "9be17748-cdad-4b87-84f2-032a8a23a11b",
        "material_id": material_id,
        "user_id": "student-id",
        "type": "text",
        "name": "Lecture notes",
        "raw_content": "Functions and variables.",
        "source_url": None,
        "storage_path": None,
        "created_at": "2026-10-15T12:00:00+00:00",
    }]

    response = client.post(
        f"/api/materials/{material_id}/sources",
        json={"type": "text", "name": "Lecture notes", "rawContent": "Functions and variables."},
    )

    assert response.status_code == 200
    assert response.json()["materialId"] == material_id
    assert response.json()["rawContent"] == "Functions and variables."
    inserted = source_table.insert.call_args.args[0]
    assert inserted["user_id"] == "student-id"
    assert inserted["material_id"] == material_id


def test_material_source_rejects_material_owned_by_another_account(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    source_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
        "material_sources": source_table,
    }[name]
    auth_login(client)
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = None

    response = client.post(
        "/api/materials/6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817/sources",
        json={"type": "text", "name": "Private", "rawContent": "Not your material."},
    )

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "material_not_found"


def test_material_file_upload_rejects_untrusted_origin(client: HttpxClient):
    response = client.post(
        "/api/materials/6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817/sources",
        files={"file": ("notes.txt", b"source", "text/plain")},
        headers={"Origin": "https://evil.example"},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "invalid_origin"


def test_source_file_text_extraction_supports_txt_and_rejects_unknown_formats():
    assert _extract_file_text("notes.txt", "text/plain", b"\xef\xbb\xbfCourse notes") == "Course notes"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Persisted PowerPoint text"
    pptx_data = BytesIO()
    presentation.save(pptx_data)
    assert _extract_file_text(
        "lecture.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        pptx_data.getvalue(),
    ) == "Persisted PowerPoint text"
    with pytest.raises(ApiError, match="Upload a TXT, PDF, or PPTX"):
        _extract_file_text("notes.doc", "application/msword", b"binary")
    with pytest.raises(ApiError, match="Private or local network"):
        _fetch_url_text("http://127.0.0.1/private")


def test_material_file_upload_stores_extracted_text_and_private_object(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    source_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
        "material_sources": source_table,
    }[name]
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Python deck",
        "type": "slides",
        "data": {"type": "slides", "content": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    source_table.insert.return_value.execute.side_effect = lambda: SimpleNamespace(
        data=[{
            **source_table.insert.call_args.args[0],
            "created_at": "2026-10-15T12:00:00+00:00",
        }]
    )
    auth_client.storage.from_.return_value.create_signed_url.return_value = {
        "signedURL": "https://example.supabase.co/signed/source",
    }

    response = client.post(
        f"/api/materials/{material_id}/sources",
        files={"file": ("notes.txt", b"Persist this source text.", "text/plain")},
    )

    assert response.status_code == 200
    assert response.json()["rawContent"] == "Persist this source text."
    assert response.json()["fileUrl"] == "https://example.supabase.co/signed/source"
    assert source_table.insert.call_args.args[0]["storage_path"].startswith(
        f"student-id/{material_id}/"
    )
    auth_client.storage.from_.return_value.upload.assert_called_once()


def test_material_file_upload_reports_missing_storage_configuration(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
    }[name]
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Python deck",
        "type": "slides",
        "data": {"type": "slides", "content": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    auth_client.storage.from_.return_value.upload.side_effect = StorageApiError(
        "Bucket not found", "BucketNotFound", 404
    )

    response = client.post(
        f"/api/materials/{material_id}/sources",
        files={"file": ("notes.txt", b"Persist this source text.", "text/plain")},
    )

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "source_bucket_missing"
    assert "20261015164500_configure_instructor_material_storage.sql" in response.json()["error"]["message"]


def test_material_file_upload_reports_storage_policy_denial(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
    }[name]
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Python deck",
        "type": "slides",
        "data": {"type": "slides", "content": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    auth_client.storage.from_.return_value.upload.side_effect = StorageApiError(
        "new row violates row-level security policy", "Unauthorized", 403
    )

    response = client.post(
        f"/api/materials/{material_id}/sources",
        files={"file": ("notes.txt", b"Persist this source text.", "text/plain")},
    )

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "source_storage_forbidden"
    assert "signed-in user JWT" in response.json()["error"]["message"]


def test_fetch_url_text_extracts_html_from_streamed_body(monkeypatch: pytest.MonkeyPatch):
    from app.routers import materials

    class FakeResponse:
        status_code = 200
        headers = {"content-type": "text/html; charset=utf-8"}
        encoding = "utf-8"

        @property
        def text(self):
            raise AssertionError("streamed HTML must be parsed from the collected body")

        def iter_bytes(self):
            yield b"<html><script>ignore me</script><body><h1>Course title</h1>"
            yield b"<p>Useful lesson text.</p></body></html>"

    class FakeHttpClient:
        def __init__(self, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def stream(self, *args, **kwargs):
            class ResponseContext:
                def __enter__(self):
                    return FakeResponse()

                def __exit__(self, *args):
                    return None

            return ResponseContext()

    monkeypatch.setattr(
        materials.socket,
        "getaddrinfo",
        lambda *args: [(None, None, None, None, ("93.184.216.34", 443))],
    )
    monkeypatch.setattr(materials, "HttpClient", FakeHttpClient)

    assert materials._fetch_url_text("https://example.com/course") == (
        "Course title\nUseful lesson text."
    )


def test_fetch_url_text_extracts_pdf_body(monkeypatch: pytest.MonkeyPatch):
    from app.routers import materials

    content = b"%PDF-1.4 test document"

    class FakeResponse:
        status_code = 200
        headers = {"content-type": "application/pdf"}
        encoding = "utf-8"

        def iter_bytes(self):
            yield content

    class FakeHttpClient:
        def __init__(self, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def stream(self, *args, **kwargs):
            class ResponseContext:
                def __enter__(self):
                    return FakeResponse()

                def __exit__(self, *args):
                    return None

            return ResponseContext()

    monkeypatch.setattr(
        materials.socket,
        "getaddrinfo",
        lambda *args: [(None, None, None, None, ("93.184.216.34", 443))],
    )
    monkeypatch.setattr(materials, "HttpClient", FakeHttpClient)
    monkeypatch.setattr(
        materials,
        "_extract_file_text",
        lambda filename, content_type, data: (
            "Extracted from a PDF source URL."
            if (filename, content_type, data) == ("source.pdf", "application/pdf", content)
            else pytest.fail("PDF URL body was not passed to the PDF text extractor")
        ),
    )

    assert materials._fetch_url_text("https://example.com/course.pdf") == (
        "Extracted from a PDF source URL."
    )


def test_material_create_and_patch_store_data_for_authenticated_owner(
    client: HttpxClient, auth_client: Mock
):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = "instructor"
    material_table = Mock()
    source_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
        "material_sources": source_table,
    }[name]
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    row = {
        "id": material_id,
        "user_id": "student-id",
        "title": "New quiz",
        "type": "quiz",
        "data": {"type": "quiz", "questions": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    material_table.insert.return_value.execute.return_value.data = [row]
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = row
    material_table.update.return_value.eq.return_value.eq.return_value.execute.return_value.data = [row]
    source_table.select.return_value.eq.return_value.eq.return_value.order.return_value.execute.return_value.data = []

    created = client.post(
        "/api/materials",
        json={
            "id": material_id,
            "title": "New quiz",
            "type": "quiz",
            "data": {"type": "quiz", "questions": []},
            "chatHistory": [],
        },
    )
    updated = client.patch(
        f"/api/materials/{material_id}",
        json={"title": "Updated quiz", "data": {"type": "quiz", "questions": []}},
    )

    assert created.status_code == 200
    assert created.json()["userId"] == "student-id"
    assert material_table.insert.call_args.args[0]["user_id"] == "student-id"
    assert updated.status_code == 200
    assert material_table.update.call_args.args[0]["title"] == "Updated quiz"


def _materials_test_tables(auth_client: Mock, role: str = "instructor"):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = role
    material_table = Mock()
    source_table = Mock()
    version_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": material_table,
        "material_sources": source_table,
        "material_versions": version_table,
    }[name]
    return material_table, source_table, version_table


def test_material_versions_require_instructor_role(client: HttpxClient) -> None:
    auth_login(client)  # default profile role is "student"
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"

    response = client.post(f"/api/materials/{material_id}/versions", json={"label": "Snapshot"})

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "instructor_required"


def test_material_versions_unknown_material_returns_404(client: HttpxClient, auth_client: Mock) -> None:
    material_table, _, _ = _materials_test_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = None

    response = client.get(f"/api/materials/{material_id}/versions")

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "material_not_found"


def test_material_version_create_list_and_restore(client: HttpxClient, auth_client: Mock) -> None:
    material_table, source_table, version_table = _materials_test_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    version_id = "1f0f1f2f-3a4b-5c6d-7e8f-9a0b1c2d3e4f"
    current_data = {
        "type": "slides",
        "content": [
            {"title": "Now", "bullets": ["b"], "visual": {"kind": "diagram", "description": "d", "steps": ["s"]}},
        ],
    }
    old_data = {
        "type": "slides",
        "content": [
            {"title": "Before", "bullets": ["b"], "visual": {"kind": "diagram", "description": "d", "steps": ["s"]}},
        ],
    }
    material_row = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Test material",
        "type": "slides",
        "data": current_data,
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    version_row = {
        "id": version_id,
        "material_id": material_id,
        "user_id": "student-id",
        "title": "Test material",
        "data": old_data,
        "label": "Before AI revision",
        "created_at": "2026-10-27T12:00:00+00:00",
    }
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = material_row
    version_table.insert.return_value.execute.return_value.data = [version_row]
    source_table.select.return_value.eq.return_value.eq.return_value.order.return_value.execute.return_value.data = []

    created = client.post(f"/api/materials/{material_id}/versions", json={"label": "Before AI revision"})

    assert created.status_code == 200
    assert created.json()["id"] == version_id
    assert created.json()["label"] == "Before AI revision"
    # The snapshot stores the material's content at the moment of capture.
    assert version_table.insert.call_args.args[0]["data"] == current_data

    version_table.select.return_value.eq.return_value.eq.return_value.order.return_value.limit.return_value.execute.return_value.data = [version_row]
    listed = client.get(f"/api/materials/{material_id}/versions")
    assert listed.status_code == 200
    assert [version["id"] for version in listed.json()] == [version_id]
    assert listed.json()[0]["data"] == old_data

    version_table.select.return_value.eq.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = version_row
    material_table.update.return_value.eq.return_value.eq.return_value.execute.return_value.data = [material_row]
    restored = client.post(f"/api/materials/{material_id}/versions/{version_id}/restore")

    assert restored.status_code == 200
    # Restore copies the version's data back into the material...
    assert material_table.update.call_args.args[0]["data"] == old_data
    # ...after first snapshotting the current state so the restore is reversible.
    assert version_table.insert.call_count == 2
    before_restore = version_table.insert.call_args_list[1].args[0]
    assert before_restore["label"] == "Before restore"
    assert before_restore["data"] == current_data


def test_material_version_restore_unknown_version_returns_404(client: HttpxClient, auth_client: Mock) -> None:
    material_table, _, version_table = _materials_test_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    version_id = "1f0f1f2f-3a4b-5c6d-7e8f-9a0b1c2d3e4f"
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Test material",
        "type": "slides",
        "data": {"type": "slides", "content": []},
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    version_table.select.return_value.eq.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = None

    response = client.post(f"/api/materials/{material_id}/versions/{version_id}/restore")

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "version_not_found"


def test_material_export_pptx_requires_instructor_role(client: HttpxClient) -> None:
    auth_login(client)  # default profile role is "student"
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"

    response = client.get(f"/api/materials/{material_id}/export/pptx")

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "instructor_required"


def test_material_export_pptx_roundtrip(client: HttpxClient, auth_client: Mock) -> None:
    material_table, _, _ = _materials_test_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    material_row = {
        "id": material_id,
        "user_id": "student-id",
        "title": "Export Deck",
        "type": "slides",
        "data": {
            "type": "slides",
            "content": [
                {
                    "title": "First slide",
                    "bullets": ["Point one", "Point two"],
                    "visual": {"kind": "diagram", "description": "flow", "steps": ["Start", "End"]},
                    "theme": {"background": "#0B192C", "textColor": "#F8FAFC", "accentColor": "#38BDF8", "cardBackground": "#1E293B"},
                },
                {"title": "Second slide", "bullets": ["Only bullet"], "visual": {"kind": "diagram", "description": "d", "steps": ["A"]}},
            ],
        },
        "chat_history": [],
        "created_at": "2026-10-15T12:00:00+00:00",
        "updated_at": "2026-10-15T12:00:00+00:00",
    }
    material_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = material_row

    response = client.get(f"/api/materials/{material_id}/export/pptx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert 'filename="Export Deck.pptx"' in response.headers["content-disposition"]
    presentation = Presentation(BytesIO(response.content))
    assert len(presentation.slides) == 2
    first_slide_text = "\n".join(
        shape.text_frame.text for shape in presentation.slides[0].shapes if shape.has_text_frame
    )
    assert "First slide" in first_slide_text
    assert "Point one" in first_slide_text


def test_build_pptx_quiz_keeps_answer_key_in_notes() -> None:
    material = {
        "title": "Sample quiz",
        "data": {
            "type": "quiz",
            "questions": [
                {
                    "id": "q1",
                    "question": "What is 2 + 2?",
                    "options": ["3", "4", "5"],
                    "correctAnswerIndex": 1,
                    "explanation": "Basic arithmetic.",
                },
            ],
        },
    }
    presentation = Presentation(BytesIO(build_pptx(material)))
    # One heading slide + one slide per question.
    assert len(presentation.slides) == 2
    notes = presentation.slides[1].notes_slide.notes_text_frame.text
    assert "Correct answer: B" in notes
    assert "Basic arithmetic." in notes
    slide_text = "\n".join(shape.text_frame.text for shape in presentation.slides[1].shapes if shape.has_text_frame)
    # The answer key must never appear on the slide itself.
    assert "Correct answer" not in slide_text
    assert "B. 4" in slide_text


def _course_tables(auth_client: Mock, role: str = "student"):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data["role"] = role
    materials_table = Mock()
    source_table = Mock()
    answers_table = Mock()
    progress_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "materials": materials_table,
        "material_sources": source_table,
        "course_quiz_answers": answers_table,
        "course_progress": progress_table,
    }[name]
    source_table.select.return_value.eq.return_value.eq.return_value.order.return_value.execute.return_value.data = []
    return materials_table, answers_table, progress_table


def _published_row(material_id: str, *, published: bool = True, data=None, title="Published course") -> dict:
    return {
        "id": material_id,
        "user_id": "instructor-id",
        "title": title,
        "type": "quiz" if (data or {}).get("type") == "quiz" else "slides",
        "data": data or {"type": "slides", "content": [{"title": "S1", "bullets": ["b"]}]},
        "chat_history": [],
        "published": published,
        "published_at": "2026-10-28T00:00:00+00:00" if published else None,
        "created_at": "2026-10-28T00:00:00+00:00",
        "updated_at": "2026-10-28T00:00:00+00:00",
    }


def test_courses_require_student_role(client: HttpxClient, auth_client: Mock) -> None:
    _course_tables(auth_client, role="instructor")
    auth_login(client)

    response = client.get("/api/courses")

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "student_required"


def test_course_list_returns_published_courses(client: HttpxClient, auth_client: Mock) -> None:
    materials_table, _, _ = _course_tables(auth_client)
    auth_login(client)
    row = _published_row("6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817", title="My quiz")
    materials_table.select.return_value.eq.return_value.order.return_value.limit.return_value.execute.return_value.data = [row]

    response = client.get("/api/courses")

    assert response.status_code == 200
    body = response.json()
    assert [item["title"] for item in body] == ["My quiz"]
    assert body[0]["publishedAt"] == "2026-10-28T00:00:00+00:00"
    # The stored material (with its answer keys) must not leak into the list.
    assert "data" not in body[0]


def test_course_get_strips_quiz_answer_keys(client: HttpxClient, auth_client: Mock) -> None:
    materials_table, _, _ = _course_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    quiz_data = {
        "type": "quiz",
        "questions": [
            {
                "id": "q1",
                "question": "What is 2 + 2?",
                "options": ["3", "4", "5"],
                "correctAnswerIndex": 1,
                "explanation": "Basic arithmetic.",
            },
        ],
    }
    materials_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = _published_row(material_id, data=quiz_data)

    response = client.get(f"/api/courses/{material_id}")

    assert response.status_code == 200
    question = response.json()["content"]["questions"][0]
    assert question["question"] == "What is 2 + 2?"
    assert question["options"] == ["3", "4", "5"]
    assert "correctAnswerIndex" not in question
    assert "explanation" not in question


def test_course_get_unpublished_returns_404(client: HttpxClient, auth_client: Mock) -> None:
    materials_table, _, _ = _course_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    materials_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = None

    response = client.get(f"/api/courses/{material_id}")

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "course_not_found"


def test_course_answer_checks_and_records(client: HttpxClient, auth_client: Mock) -> None:
    materials_table, answers_table, _ = _course_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    quiz_data = {
        "type": "quiz",
        "questions": [
            {
                "id": "q1",
                "question": "What is 2 + 2?",
                "options": ["3", "4", "5"],
                "correctAnswerIndex": 1,
                "explanation": "Basic arithmetic.",
            },
        ],
    }
    materials_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = _published_row(material_id, data=quiz_data)
    answers_table.insert.return_value.execute.return_value.data = [{"id": "a1"}]

    wrong = client.post(f"/api/courses/{material_id}/answers", json={"question_id": "q1", "selected_index": 0})
    assert wrong.status_code == 200
    assert wrong.json() == {"correct": False, "correct_index": 1, "explanation": "Basic arithmetic."}
    assert answers_table.insert.call_args.args[0]["correct"] is False

    right = client.post(f"/api/courses/{material_id}/answers", json={"question_id": "q1", "selected_index": 1})
    assert right.status_code == 200
    assert right.json()["correct"] is True
    assert answers_table.insert.call_args.args[0]["user_id"] == "student-id"

    unknown = client.post(f"/api/courses/{material_id}/answers", json={"question_id": "nope", "selected_index": 0})
    assert unknown.status_code == 404
    assert unknown.json()["error"]["code"] == "question_not_found"


def test_course_progress_start_and_complete(client: HttpxClient, auth_client: Mock) -> None:
    materials_table, _, progress_table = _course_tables(auth_client)
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    materials_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = _published_row(material_id)

    started = client.post(f"/api/courses/{material_id}/progress", json={"event": "start"})
    assert started.status_code == 204
    start_call = progress_table.upsert.call_args
    assert start_call.args[0]["user_id"] == "student-id"
    assert start_call.kwargs["ignore_duplicates"] is True

    completed = client.post(f"/api/courses/{material_id}/progress", json={"event": "complete"})
    assert completed.status_code == 204
    complete_call = progress_table.upsert.call_args
    assert "completed_at" in complete_call.args[0]


def test_publish_material_toggles_flag(client: HttpxClient, auth_client: Mock) -> None:
    materials_table, _, _ = _course_tables(auth_client, role="instructor")
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    materials_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = _published_row(material_id, published=False)
    materials_table.update.return_value.eq.return_value.eq.return_value.execute.return_value.data = [_published_row(material_id, published=True)]

    published = client.post(f"/api/materials/{material_id}/publish", json={"published": True})
    assert published.status_code == 200
    assert published.json()["published"] is True
    update_args = materials_table.update.call_args.args[0]
    assert update_args["published"] is True
    assert update_args["published_at"] is not None

    captured = {}
    materials_table.update.side_effect = lambda changes: (captured.update(changes) or materials_table.update.return_value)
    unpublished = client.post(f"/api/materials/{material_id}/publish", json={"published": False})
    assert unpublished.status_code == 200
    assert captured["published"] is False
    assert captured["published_at"] is None


def test_publish_material_requires_instructor(client: HttpxClient, auth_client: Mock) -> None:
    _course_tables(auth_client, role="student")
    auth_login(client)
    material_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"

    response = client.post(f"/api/materials/{material_id}/publish", json={"published": True})

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "instructor_required"


def test_roster_csv_parsing() -> None:
    from app.identity.errors import ApiError
    from app.routers.rosters import _parse_roster_csv

    rows = _parse_roster_csv("Email,Name\na@x.edu,Alice\nb@x.edu,Bob\n")
    assert rows == [
        {"email": "a@x.edu", "name": "Alice"},
        {"email": "b@x.edu", "name": "Bob"},
    ]
    # Duplicate emails collapse; missing names fall back to the local part.
    rows = _parse_roster_csv("name,email\nAlice,a@x.edu\n,a@x.edu\n")
    assert rows == [{"email": "a@x.edu", "name": "Alice"}]

    with pytest.raises(ApiError) as excinfo:
        _parse_roster_csv("name\nAlice\n")
    assert excinfo.value.code == "roster_csv_invalid"

    with pytest.raises(ApiError) as excinfo:
        _parse_roster_csv("name,email\nAlice,nah\n")
    assert excinfo.value.code == "roster_csv_invalid"

    with pytest.raises(ApiError) as excinfo:
        _parse_roster_csv("name,email\n")
    assert excinfo.value.code == "roster_csv_empty"


def _roster_tables(auth_client: Mock, role: str = "instructor"):
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data[
        "role"
    ] = role
    rosters_table = Mock()
    students_table = Mock()
    progress_table = Mock()
    answers_table = Mock()
    materials_table = Mock()
    auth_client.table.side_effect = lambda name: {
        "profiles": profile_table,
        "rosters": rosters_table,
        "roster_students": students_table,
        "course_progress": progress_table,
        "course_quiz_answers": answers_table,
        "materials": materials_table,
    }[name]
    return rosters_table, students_table, progress_table, answers_table, materials_table


def test_roster_import_requires_instructor(client: HttpxClient, auth_client: Mock) -> None:
    _roster_tables(auth_client, role="student")
    auth_login(client)

    response = client.post(
        "/api/rosters/import",
        files={"file": ("roster.csv", b"name,email\na@x.edu,Alice\n", "text/csv")},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "instructor_required"


def test_roster_import_invalid_csv(client: HttpxClient, auth_client: Mock) -> None:
    rosters_table, *_ = _roster_tables(auth_client)
    auth_login(client)

    response = client.post(
        "/api/rosters/import",
        files={"file": ("roster.csv", b"name\nAlice\n", "text/csv")},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "roster_csv_invalid"
    rosters_table.insert.assert_not_called()


def test_roster_import_creates_students_and_matches(client: HttpxClient, auth_client: Mock) -> None:
    rosters_table, students_table, *_ = _roster_tables(auth_client)
    auth_login(client)
    rosters_table.insert.return_value.execute.return_value.data = [
        {"id": "r-1", "instructor_id": "student-id", "created_at": "2026-10-29T00:00:00+00:00"},
    ]
    students_table.upsert.return_value.execute.return_value.data = []
    profile_table = auth_client.table.return_value
    profile_table.select.return_value.in_.return_value.limit.return_value.execute.return_value.data = [
        {"id": "student-id", "email": "a@x.edu"},
    ]

    response = client.post(
        "/api/rosters/import",
        files={"file": ("fall.csv", b"name,email\nAlice,a@x.edu\nBob,b@x.edu\n", "text/csv")},
        data={"name": "Fall 2026"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["imported"] == 2
    assert body["matched"] == 1
    assert body["roster"]["name"] == "Fall 2026"
    upserted = students_table.upsert.call_args.args[0]
    assert [row["email"] for row in upserted] == ["a@x.edu", "b@x.edu"]
    assert upserted[0]["user_id"] == "student-id"
    assert upserted[1]["user_id"] is None
    assert students_table.upsert.call_args.kwargs["on_conflict"] == "roster_id,lower(email)"


def test_roster_analytics_aggregates(client: HttpxClient, auth_client: Mock) -> None:
    rosters_table, students_table, progress_table, answers_table, materials_table = _roster_tables(auth_client)
    auth_login(client)
    roster_id = "6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817"
    rosters_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = {
        "id": roster_id,
        "instructor_id": "student-id",
        "name": "Fall 2026",
        "created_at": "2026-10-29T00:00:00+00:00",
    }
    students_table.select.return_value.eq.return_value.order.return_value.execute.return_value.data = [
        {"id": "s1", "roster_id": roster_id, "user_id": "student-id", "email": "a@x.edu", "name": "Alice", "imported_at": "2026-10-29T00:00:00+00:00"},
        {"id": "s2", "roster_id": roster_id, "user_id": None, "email": "b@x.edu", "name": "Bob", "imported_at": "2026-10-29T00:00:00+00:00"},
    ]
    progress_table.select.return_value.in_.return_value.execute.return_value.data = [
        {"user_id": "student-id", "material_id": "c1", "started_at": "2026-10-29T00:00:00+00:00", "last_active_at": "2026-10-29T01:00:00+00:00", "completed_at": "2026-10-29T02:00:00+00:00"},
    ]
    answers_table.select.return_value.in_.return_value.execute.return_value.data = [
        {"user_id": "student-id", "material_id": "c1", "correct": True},
        {"user_id": "student-id", "material_id": "c1", "correct": False},
    ]
    materials_table.select.return_value.in_.return_value.limit.return_value.execute.return_value.data = [
        {"id": "c1", "title": "Thermo"},
    ]

    response = client.get(f"/api/rosters/{roster_id}")

    assert response.status_code == 200
    body = response.json()
    analytics = body["analytics"]
    assert analytics["enrolled"] == 2
    assert analytics["matched"] == 1
    assert analytics["active"] == 1
    assert analytics["participationRate"] == 1.0
    assert analytics["quizAverage"] == 50.0
    assert analytics["courses"][0]["title"] == "Thermo"
    assert analytics["courses"][0]["started"] == 1
    assert analytics["courses"][0]["completed"] == 1
    alice = body["students"][0]
    assert alice["matched"] is True
    assert alice["coursesCompleted"] == 1
    assert alice["quizAverage"] == 50.0
    assert body["students"][1]["matched"] is False


def test_roster_detail_unknown_returns_404(client: HttpxClient, auth_client: Mock) -> None:
    rosters_table, *_ = _roster_tables(auth_client)
    auth_login(client)
    rosters_table.select.return_value.eq.return_value.eq.return_value.maybe_single.return_value.execute.return_value.data = None

    response = client.get("/api/rosters/6b6bcbf4-9beb-48fd-a4c6-1f3ed71eb817")

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "roster_not_found"


def test_instructor_course_generation_rejects_mismatched_existing_content(client: HttpxClient):
    response = client.post(
        "/api/instructor/generate",
        json={
            "artifact_type": "slides",
            "prompt": "Revise the resource.",
            "existing_content": {
                "type": "quiz",
                "questions": [{
                    "id": "q1",
                    "question": "Question?",
                    "options": ["A", "B"],
                    "correctAnswerIndex": 0,
                    "explanation": "Because.",
                }],
            },
        },
    )
    assert response.status_code == 422


def test_course_generator_uses_configured_key_and_validates_slides(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    http_client = Mock()
    http_client.post.return_value.status_code = 200
    http_client.post.return_value.json.return_value = {
        "choices": [{
            "message": {
                "content": '{"type":"slides","content":[{"title":"Python Basics","bullets":["Variables"],"visual":{"kind":"diagram","description":"A variable flows into a function","steps":["Variable","Function"]}}]}',
            },
        }],
    }

    result = generate_course_content("Teach Python basics.", "Variables", http_client)

    assert result.type == "slides"
    assert result.content[0].title == "Python Basics"
    assert http_client.post.call_args.kwargs["headers"]["Authorization"] == "Bearer test-openrouter-key"
    assert http_client.post.call_args.kwargs["json"]["messages"][1]["content"].endswith("Variables")
    assert result.content[0].visual.kind == "diagram"


def test_course_generator_preserves_requested_slide_theme(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    http_client = Mock()
    http_client.post.return_value.status_code = 200
    http_client.post.return_value.json.return_value = {
        "choices": [{
            "message": {
                "content": (
                    '{"type":"slides","content":[{"title":"Python Basics","bullets":["Variables"],'
                    '"visual":{"kind":"diagram","description":"Variable flow","steps":["Input","Output"]},'
                    '"theme":{"background":"#0B192C","textColor":"#F8FAFC",'
                    '"accentColor":"#38BDF8","cardBackground":"#1E293B"}}]}'
                ),
            },
        }],
    }

    result = generate_course_content(
        "Create a slide deck with a dark blue theme.",
        "",
        http_client,
    )

    assert result.content[0].theme is not None
    assert result.content[0].theme.background == "#0B192C"
    assert result.content[0].theme.text_color == "#F8FAFC"
    assert result.content[0].theme.card_background == "#1E293B"
    system_prompt = http_client.post.call_args.kwargs["json"]["messages"][0]["content"]
    assert "Extract any requested aesthetic or color palette" in system_prompt


def test_course_generator_requires_openrouter_key(monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)

    with pytest.raises(CourseGeneratorNotConfiguredError):
        generate_course_content("Teach Python basics.", "", Mock())


def test_course_generator_rejects_invalid_slide_response(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    http_client = Mock()
    http_client.post.return_value.status_code = 200
    http_client.post.return_value.json.return_value = {
        "choices": [{"message": {"content": '{"type":"slides","content":[]}'}}],
    }

    with pytest.raises(CourseGeneratorUnavailableError):
        generate_course_content("Teach Python basics.", "", http_client)


def test_course_generator_includes_existing_slides_and_revision_prompt(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    http_client = Mock()
    http_client.post.return_value.status_code = 200
    http_client.post.return_value.json.return_value = {
        "choices": [{
            "message": {
                "content": '{"type":"slides","content":[{"title":"Revised","bullets":["One"],"visual":{"kind":"diagram","description":"Flow","steps":["A","B"]}}]}',
            },
        }],
    }
    existing = GeneratedCourseContent.model_validate({
        "type": "slides",
        "content": [{
            "title": "Original",
            "bullets": ["Existing bullet"],
            "visual": {"kind": "diagram", "description": "Flow", "steps": ["A", "B"]},
        }],
    })

    generate_course_content(
        "",
        "",
        http_client,
        existing_content=existing,
        revision_prompt="Add a practical example.",
    )

    request_content = http_client.post.call_args.kwargs["json"]["messages"][1]["content"]
    assert "Add a practical example." in request_content
    assert "Existing bullet" in request_content


def test_course_generator_builds_validated_quiz(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    http_client = Mock()
    http_client.post.return_value.status_code = 200
    http_client.post.return_value.json.return_value = {
        "choices": [{
            "message": {
                "content": (
                    '{"type":"quiz","questions":[{"id":"q1","question":"Which is an integer?",'
                    '"options":["int","string","float","boolean"],"correctAnswerIndex":0,'
                    '"explanation":"int stores whole numbers."}]}'
                ),
            },
        }],
    }

    result = generate_course_content("Create a quiz.", "", http_client, artifact_type="quiz")

    assert result.type == "quiz"
    assert result.questions[0].correct_answer_index == 0
    assert http_client.post.call_args.kwargs["json"]["messages"][0]["content"].count("multiple-choice") == 1


def test_course_generator_rejects_quiz_answer_index_out_of_range(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-openrouter-key")
    http_client = Mock()
    http_client.post.return_value.status_code = 200
    http_client.post.return_value.json.return_value = {
        "choices": [{
            "message": {
                "content": (
                    '{"type":"quiz","questions":[{"id":"q1","question":"Question?",'
                    '"options":["A","B","C","D"],"correctAnswerIndex":4,"explanation":"Because."}]}'
                ),
            },
        }],
    }

    with pytest.raises(CourseGeneratorUnavailableError):
        generate_course_content("Create a quiz.", "", http_client, artifact_type="quiz")


def test_judge_parse_verdict_is_defensive():
    verdict = parse_verdict('{"score": 80, "rating": "sharp", "identified": ["a"], "missed": [], "feedback": "Good."}')
    assert (verdict.score, verdict.rating, verdict.identified, verdict.missed, verdict.feedback) == (
        80, "sharp", ["a"], [], "Good.",
    )
    padded = parse_verdict('Sure! Here is the verdict: {"score": 40, "feedback": "ok"}')
    assert padded.score == 40 and padded.rating == "novice" and padded.feedback == "ok"
    with pytest.raises(JudgeUnavailableError):
        parse_verdict("no json at all")
    with pytest.raises(JudgeUnavailableError):
        parse_verdict("[1, 2, 3]")
